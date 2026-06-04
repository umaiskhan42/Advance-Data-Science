"""
Generate the 4-page Final Project Proposal / Report as a Word .docx, following a
standard academic project-proposal template and populated with this project's
real content, tables, and figures.

Output: Final_Project_Proposal_Umais_Khan.docx  (project root)
"""
from __future__ import annotations
import os
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config as C

FIG = C.RES_FIG
OUT = os.path.join(C.ROOT, "Final_Project_Proposal_Umais_Khan.docx")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def setup(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10)
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.05


def title_block(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Zero-Shot Microfauna Anomaly Detection with Frozen DINOv3 "
                  "on MD-AS-2025")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Final Project Proposal & Report  |  Applied Data Science")
    r.italic = True; r.font.size = Pt(10.5)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Umais Khan  ·  Student ID 2025311037  ·  Dept. of Environmental "
              "Science & Engineering\n").font.size = Pt(9.5)
    p.add_run("Supervisor: Prof. Won Hee Lee  ·  EMSEL (Environmental "
              "Management & Systems Engineering Lab)").font.size = Pt(9.5)
    doc.add_paragraph()


def heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12)
    r.font.color.rgb = NAVY
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def figure(doc, path, width, caption):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(8.5)
    c.paragraph_format.space_after = Pt(5)


def table(doc, header, rows, caption=None, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        t.style = "Light Grid Accent 1"
    except Exception:
        t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
        for r in cell.paragraphs[0].runs:
            r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for para in cells[i].paragraphs:
                for r in para.runs:
                    r.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    if caption:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption); r.italic = True; r.font.size = Pt(8.5)
        c.paragraph_format.space_after = Pt(5)
    return t


def build():
    doc = Document()
    setup(doc)
    title_block(doc)

    # ---- Abstract -------------------------------------------------------- #
    heading(doc, "Abstract")
    body(doc,
         "Sludge degradation in activated-sludge wastewater treatment is signalled "
         "by a shift in the microfauna community from healthy ciliates and rotifers "
         "to stressed nematodes and testate amoebae, yet labelled degradation "
         "events are rare and manual microscopy is slow. We treat the problem as "
         "zero-shot anomaly detection: a frozen DINOv3 vision-foundation model "
         "supplies patch embeddings, a PatchCore memory bank stores only healthy "
         "microfauna, and test images are scored by their nearest-neighbour distance "
         "to that healthy manifold. With no task supervision and no fine-tuning the "
         "method reaches 0.79 image-level AUROC, far above a random-init ViT (0.53) "
         "and an ImageNet ResNet-50 (0.59). We additionally use the same frozen "
         "backbone for object detection, reproduce the dataset paper's figure suite, "
         "and quantify, in a control-loop simulation, the downstream value of the "
         "detector for plant operation. The work is a direct extension of EMSEL's "
         "memory/prototype-based anomaly-detection lineage from sensor data into "
         "computer vision.")

    # ---- 1. Introduction ------------------------------------------------- #
    heading(doc, "1. Introduction & Motivation")
    body(doc,
         "Aeration consumes 40-75% of the electricity used at activated-sludge "
         "plants, and the biological condition of the sludge is the primary "
         "determinant of treatment efficiency and aeration demand. Microfauna are a "
         "well-established bio-indicator of that condition (the Sludge Biotic Index, "
         "Madoni 1994): a community dominated by ciliates/rotifers indicates a "
         "well-performing plant, whereas a shift toward nematodes and testate "
         "amoebae signals stress. Supervised vision is a poor fit because labelled "
         "degradation events are rare and visual heterogeneity across plants and "
         "seasons is high. We therefore pose degradation as anomaly detection: model "
         "the 'normal' (healthy) microfauna and flag deviations, requiring labels for "
         "neither the anomalies nor the model itself.")

    # ---- 2. Relationship to prior lab work ------------------------------- #
    heading(doc, "2. Relationship to Prior Lab Work (an extension)")
    body(doc,
         "This project is positioned as the next step in EMSEL's anomaly-detection "
         "research rather than a separate line. The lab's 'model-the-normal, "
         "flag-the-deviation' lineage runs from KPCA/ICA process monitoring (Lee & "
         "Yoo, 2004), through deep variational-residual autoencoders for WWTP sensor "
         "fault detection (Ba-Alawi et al., 2022), to memory-augmented autoencoders "
         "with learned normal prototypes (Ghorbani, Tariq & Yoo, 2025). Our PatchCore "
         "memory bank of healthy DINOv3 patch embeddings is the direct image-domain "
         "analogue of that 2025 normal-prototype memory. Only two things change: the "
         "modality (sensor time-series to microscopy images) and the feature source "
         "(a trained autoencoder to a frozen, label-free foundation model).")

    # ---- 3. Dataset ------------------------------------------------------ #
    heading(doc, "3. Dataset: MD-AS-2025")
    body(doc,
         "MD-AS-2025 (Pang et al., Scientific Data 2025) provides 14,257 cropped "
         "single-object microscopy images across five taxa plus 4,000 annotated "
         "full-frame images. We map taxa to the anomaly-detection roles below; the "
         "memory bank is built from healthy taxa only, and stressed taxa are never "
         "seen during model construction.")
    table(doc,
          ["Role", "Taxa", "Label", "Cropped images"],
          [["Healthy (normal)", "Hypotrichida; Digononta; Monogononta", "0",
            "556; 4,780; 1,207"],
           ["Stressed (anomaly)", "Arcellinida; Nematoda", "1", "7,007; 707"]],
          "Table 1. MD-AS-2025 taxa and their anomaly-detection roles.")
    figure(doc, os.path.join(FIG, "dataset_distribution.png"), 5.6,
           "Figure 1. Per-taxon image counts and the sampled train/test split "
           "(healthy = memory bank + healthy test; stressed = test only).")

    # ---- 4. Methodology -------------------------------------------------- #
    heading(doc, "4. Methodology")
    body(doc,
         "Backbone. DINOv3 ViT-S/16 (timm vit_small_patch16_dinov3.lvd1689m, "
         "LVD-1689M pretraining) is used frozen - no gradient updates, no labels. "
         "For each image we extract the 14x14 grid of patch tokens plus the CLS "
         "token at 224 px.")
    body(doc,
         "Anomaly detection (PatchCore). Patch features are locally aggregated "
         "(3x3), L2-normalised, and pooled across the healthy training images; a "
         "greedy k-center coreset (5,000 vectors) forms the memory bank. Each test "
         "patch is scored by its nearest-neighbour cosine distance to the bank, and "
         "the image score is the maximum patch distance. Thresholds are set without "
         "labels (95th percentile of healthy scores) and at Youden-J. We compare "
         "against two baselines under the identical pipeline: a random-init ViT "
         "(isolates the value of pretraining) and an ImageNet ResNet-50 (the original "
         "PatchCore CNN backbone).")
    body(doc,
         "Object detection. To keep the project on one frozen backbone, the same "
         "DINOv3 features (cached once at 448 px, 28x28 grid) feed a lightweight "
         "anchor-free head (objectness + 5-class + box) trained on the 3,600/400 "
         "split used by the dataset paper. Only the head is trained.")

    # ---- 5. Results ------------------------------------------------------ #
    heading(doc, "5. Results")
    body(doc,
         "H1 - DINOv3 beats the baselines. On 2,154 test images (647 healthy, 1,507 "
         "stressed) the frozen DINOv3 PatchCore reaches 0.79 AUROC, with the "
         "random-init ViT near chance and the CNN well behind - the signal comes from "
         "self-supervised pretraining, not the architecture. Per-failure-mode AUROC "
         "is balanced (Arcellinida 0.788, Nematoda 0.793).")
    table(doc,
          ["Backbone / method", "AUROC", "AP", "best F1"],
          [["DINOv3 (frozen) PatchCore", "0.791", "0.899", "0.761"],
           ["DINOv3 (frozen) CLS-kNN", "0.781", "0.894", "-"],
           ["ResNet-50 (ImageNet) PatchCore", "0.591", "0.783", "0.565"],
           ["Random-init ViT PatchCore", "0.526", "0.736", "0.442"]],
          "Table 2. Image-level anomaly detection (healthy vs stressed).")
    figure(doc, os.path.join(FIG, "auroc_comparison.png"), 3.6,
           "Figure 2. AUROC by backbone (H1).")
    body(doc,
         "H2/H3 - Attention & feature space. DINOv3 self-attention overlays "
         "concentrate on the discriminative morphology (the elongated nematode body, "
         "the discoid testate shell), and t-SNE/UMAP of CLS embeddings separate "
         "stressed from healthy taxa - both produced by the pipeline (results/).")
    body(doc,
         "Object detection vs the dataset paper. With a frozen backbone and a tiny "
         "single-scale head (no fine-tuning) we obtain mAP@0.5 = 0.282 versus the "
         "paper's fully-trained YOLOv5/v8 at 0.895. The confusion matrix shows that "
         "when the detector fires it assigns the correct taxon; errors are dominated "
         "by missed objects (recall gap), and the thin Nematoda is hardest. This "
         "cleanly delineates where frozen features help (image-level judgements) and "
         "where end-to-end training is still needed (precise localisation).")
    table(doc,
          ["Taxon", "Frozen DINOv3 (ours) AP@0.5", "Paper (YOLOv5/8) mAP@0.5"],
          [["Arcellinida", "0.444", "0.958"],
           ["Hypotrichida", "0.539", "0.991"],
           ["Digononta", "0.212", "0.916"],
           ["Monogononta", "0.151", "0.834"],
           ["Nematoda", "0.065", "0.778"],
           ["Overall", "0.282", "0.895"]],
          "Table 3. Detection AP@0.5 per taxon vs the MD-AS-2025 paper.")

    # ---- 6. WWTP impact -------------------------------------------------- #
    heading(doc, "6. Anomaly Detection as a WWTP Component (with vs without)")
    body(doc,
         "Because microfauna stress precedes effluent failure, the detector acts as "
         "a leading indicator. A control-loop scenario simulation - parameterised by "
         "the detector's real operating point (TPR 0.68, FPR 0.24) - compares "
         "reactive operation (act only when the effluent violates the limit) with "
         "proactive operation (act on the detector alarm).")
    table(doc,
          ["KPI", "Without AD", "With AD", "Effect"],
          [["Action day (early warning)", "34", "24", "10 days earlier"],
           ["Effluent violation days", "4", "0", "full compliance"],
           ["Days in bulking (SVI>150)", "12", "0", "bulking avoided"],
           ["Aeration energy (kWh)", "99,196", "96,186", "-3,010 (~3%)"],
           ["Total operating cost ($)", "13,904", "11,542", "-2,362"]],
          "Table 4. Effect of the anomaly-detection component over one episode.")
    figure(doc, os.path.join(FIG, "wwtp_with_vs_without_ad.png"), 5.8,
           "Figure 3. Sludge health, detector signal, effluent quality and "
           "cumulative aeration energy - with vs without the detector.")

    # ---- 7. Conclusion --------------------------------------------------- #
    heading(doc, "7. Conclusion & Future Work")
    body(doc,
         "A frozen, label-free vision foundation model detects stress-indicating "
         "microfauna at ~0.79 AUROC with zero task supervision, confirming H1-H3 and "
         "supporting H4: self-supervised vision can be a low-cost, label-efficient "
         "front-end for WWTP / aeration-energy management. Embedded as a component it "
         "delivers earlier warning, effluent compliance, avoided bulking and lower "
         "operating cost. The detector is the image-domain continuation of EMSEL's "
         "normal-prototype anomaly detection. Future work: multi-scale features and "
         "light backbone fine-tuning to close the detection gap, temporal aggregation "
         "of plant-scale microfauna signals, and coupling the detector output to the "
         "lab's reinforcement-learning aeration controllers.")

    # ---- References ------------------------------------------------------ #
    heading(doc, "References")
    refs = [
        "Pang, Z. et al. (2025). Image Dataset of Microfauna in Activated Sludge "
        "(MD-AS-2025). Scientific Data 12:1972.",
        "Simeoni, O. et al. (2025). DINOv3. arXiv:2508.10104.",
        "Roth, K. et al. (2022). Towards Total Recall in Industrial Anomaly "
        "Detection (PatchCore). CVPR.",
        "Oquab, M. et al. (2023). DINOv2. arXiv:2304.07193.  Caron, M. et al. "
        "(2021). Emerging Properties in Self-Supervised ViTs (DINO). ICCV.",
        "Ghorbani, V., Tariq, S. & Yoo, C. (2025). Validation of subway IAQ data "
        "using memory-augmented autoencoders with learned normal prototypes. "
        "Korean J. Chem. Eng. 42:2231-2252.",
        "Ba-Alawi, A. H., Loy-Benitez, J., Kim, S. & Yoo, C. (2022). Sensor "
        "self-validation for WWTPs via deep variational residual autoencoders. "
        "Chemosphere 288:132647.",
        "Madoni, P. (1994). A sludge biotic index (SBI) for activated-sludge plants. "
        "Water Research 28(1):67-75.",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"[{i}] {r}"); run.font.size = Pt(8.5)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
